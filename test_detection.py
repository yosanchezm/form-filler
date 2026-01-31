"""
Script para probar la detección de campos en documentos DOCX.
Útil para diagnóstico cuando hay campos que no se detectan.
"""

import sys
import io
from pathlib import Path

from docx import Document

# Importar utilidades del proyecto
from utils.docx_utils import (
    extract_bracket_tokens_from_doc,
    iter_all_paragraphs,
    detect_formatted_placeholders,
    get_paragraph_full_text,
    has_shading,
    extract_shaded_text
)
from utils.text_utils import (
    first_underscore_span,
    find_all_underscore_spans,
    extract_field_context,
    detect_date_field
)


def analyze_document(doc_path: str):
    """
    Analiza un documento DOCX y muestra todos los campos detectados.
    """
    print(f"\n{'='*60}")
    print(f"Analizando: {doc_path}")
    print(f"{'='*60}\n")
    
    doc = Document(doc_path)
    
    # 1. Tokens bracket [...]
    print("📌 TOKENS BRACKET [...]")
    print("-" * 40)
    tokens = extract_bracket_tokens_from_doc(doc)
    if tokens:
        for i, tok in enumerate(tokens, 1):
            print(f"  {i}. {tok}")
    else:
        print("  (No se encontraron tokens bracket)")
    print()
    
    # 2. Párrafos con underscores
    print("📝 CAMPOS CON UNDERSCORES (____)")
    print("-" * 40)
    underscore_count = 0
    for where, p in iter_all_paragraphs(doc):
        full = get_paragraph_full_text(p)
        spans = find_all_underscore_spans(full, min_length=4)
        if spans:
            underscore_count += len(spans)
            print(f"  📍 {where}")
            context = extract_field_context(full)
            print(f"     Label: {context.get('label', '(sin label)')}")
            print(f"     Spans: {len(spans)} campos de underscore")
            
            # Verificar si es campo de fecha
            if detect_date_field(full):
                print(f"     ⚠️ Campo de FECHA detectado")
            print()
    
    if underscore_count == 0:
        print("  (No se encontraron campos con underscores)")
    print(f"  Total: {underscore_count} campos de underscore\n")
    
    # 3. Placeholders formateados (negrita, subrayado, fondo gris)
    print("🎨 PLACEHOLDERS FORMATEADOS")
    print("-" * 40)
    formatted_count = 0
    for where, p in iter_all_paragraphs(doc):
        placeholders = detect_formatted_placeholders(p)
        if placeholders:
            for ph in placeholders:
                formatted_count += 1
                print(f"  📍 {where}")
                print(f"     Texto: {ph['text'][:100]}{'...' if len(ph['text']) > 100 else ''}")
                print(f"     Formato: {', '.join(ph['format'])}")
                print()
    
    if formatted_count == 0:
        print("  (No se encontraron placeholders formateados)")
    print(f"  Total: {formatted_count} placeholders formateados\n")
    
    # 4. Texto con fondo gris (shading)
    print("🔳 TEXTO CON FONDO GRIS/SOMBREADO")
    print("-" * 40)
    shaded_count = 0
    for where, p in iter_all_paragraphs(doc):
        shaded = extract_shaded_text(p)
        if shaded:
            for text in shaded:
                shaded_count += 1
                print(f"  📍 {where}")
                print(f"     Texto: {text[:100]}{'...' if len(text) > 100 else ''}")
                print()
    
    if shaded_count == 0:
        print("  (No se encontró texto con sombreado)")
    print(f"  Total: {shaded_count} bloques sombreados\n")
    
    # 5. Análisis por runs (para debug de formateo)
    print("🔍 ANÁLISIS DETALLADO DE RUNS (primeros 10 párrafos)")
    print("-" * 40)
    count = 0
    for where, p in iter_all_paragraphs(doc):
        if count >= 10:
            break
        full = get_paragraph_full_text(p)
        if not full.strip():
            continue
        
        count += 1
        print(f"\n  📍 {where}")
        print(f"     Texto completo: {full[:150]}{'...' if len(full) > 150 else ''}")
        
        if p.runs:
            print(f"     Runs ({len(p.runs)}):")
            for i, run in enumerate(p.runs):
                if run.text:
                    formats = []
                    if run.bold:
                        formats.append("bold")
                    if run.underline:
                        formats.append("underline")
                    if has_shading(run):
                        formats.append("shading")
                    
                    format_str = f" [{', '.join(formats)}]" if formats else ""
                    print(f"        {i}: '{run.text[:50]}'{format_str}")
    
    # Resumen
    print(f"\n{'='*60}")
    print("RESUMEN")
    print(f"{'='*60}")
    print(f"  • Tokens bracket [...]: {len(tokens)}")
    print(f"  • Campos underscore: {underscore_count}")
    print(f"  • Placeholders formateados: {formatted_count}")
    print(f"  • Bloques con sombreado: {shaded_count}")
    print(f"  • Total campos detectados: {len(tokens) + underscore_count + formatted_count}")


def main():
    if len(sys.argv) < 2:
        print("Uso: python test_detection.py <documento.docx>")
        print("\nEjemplo:")
        print("  python test_detection.py mi_formulario.docx")
        sys.exit(1)
    
    doc_path = sys.argv[1]
    
    if not Path(doc_path).exists():
        print(f"Error: No se encontró el archivo: {doc_path}")
        sys.exit(1)
    
    analyze_document(doc_path)


if __name__ == "__main__":
    main()
